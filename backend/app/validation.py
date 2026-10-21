import csv
import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd

from .rules import SHIPPED, STATUSES, TERMINAL
from .timeutils import DATE_ONLY, IST, local_day, normalize_date, timestamp

HEADERS = {
    "Order ID": "id", "Order date": "order_date", "Customer name": "customer_name",
    "Phone number": "phone", "Order value": "order_value",
    "Payment type (Prepaid or COD)": "payment_type", "Order status": "status",
    "Dispatch date": "dispatch_date", "Expected delivery date": "expected_delivery_date",
    "Assigned employee": "assigned_employee",
}
OPTIONAL = {"confirmed_at", "delivery_attempt_at"}
ALIASES = {re.sub(r"[^a-z0-9]", "", k.lower()): v for k, v in HEADERS.items()}
ALIASES.update({re.sub(r"[^a-z0-9]", "", v): v for v in list(HEADERS.values()) + list(OPTIONAL)})
ALIASES.update({"paymenttype": "payment_type", "phonenumber": "phone", "orderid": "id", "orderstatus": "status"})


class FileError(ValueError):
    pass


def clean(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if isinstance(value, datetime):
        # Without cell precision metadata, midnight is conservatively a date.
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    return str(value).strip()


def read_file(filename, content):
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".csv":
            decoded = content.decode("utf-8-sig")
            parsed = list(csv.reader(io.StringIO(decoded), strict=True))
            if not parsed:
                raise FileError("The file is empty.")
            width = len(parsed[0])
            if any(len(r) != width for r in parsed[1:] if any(r)):
                raise FileError("CSV rows must contain the same number of cells as the header. Quote commas inside cells.")
            source_columns = parsed[0]
            frame = pd.read_csv(io.StringIO(decoded), dtype=str, keep_default_na=False, skip_blank_lines=False)
        elif ext in {".xlsx", ".xls"}:
            header_frame = pd.read_excel(io.BytesIO(content), header=None, nrows=1, dtype=object, keep_default_na=False)
            if header_frame.empty:
                raise FileError("The workbook is empty.")
            source_columns = header_frame.iloc[0].tolist()
            frame = pd.read_excel(io.BytesIO(content), dtype=object, keep_default_na=False)
            if ext == ".xlsx":
                from openpyxl import load_workbook
                sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
                for ri, cells in enumerate(sheet.iter_rows(min_row=2)):
                    if ri >= len(frame):
                        break
                    for ci, cell in enumerate(cells):
                        if ci >= len(frame.columns):
                            break
                        if isinstance(cell.value, datetime):
                            fmt = re.sub(r'"[^\"]*"|\[[^\]]*\]', "", cell.number_format.lower())
                            frame.iat[ri, ci] = cell.value.isoformat() if ("h" in fmt or "s" in fmt) else cell.value.date().isoformat()
                sheet.parent.close()
        else:
            raise FileError("Choose a UTF-8 .csv, .xlsx, or .xls file.")
    except FileError:
        raise
    except Exception as exc:
        raise FileError(f"File could not be read: {type(exc).__name__}. Check encoding, format, and workbook protection.") from exc
    if len(frame) > 50000:
        raise FileError("Limit each upload to 50,000 rows.")
    columns = [ALIASES.get(re.sub(r"[^a-z0-9]", "", str(c).lower()), str(c).strip()) for c in source_columns]
    if len(columns) != len(set(columns)):
        raise FileError("Duplicate column headers are ambiguous. Give each field one column.")
    frame.columns = columns
    missing = set(HEADERS.values()) - set(columns)
    # Missing columns become per-row correction issues; readable rows are retained.
    for col in missing:
        frame[col] = ""
    return [{str(k): clean(v) for k, v in row.items()} for row in frame.to_dict("records")], sorted(missing)


def validate(raw, config, now):
    row = {k: clean(v) for k, v in raw.items()}
    errors = []

    def error(field, why, correction):
        errors.append({"field": field, "explanation": why, "recommended_correction": correction})

    for field in ("id", "order_date", "customer_name", "phone", "order_value", "payment_type", "status", "assigned_employee"):
        if not row.get(field):
            error(field, "Required field is missing.", f"Supply {field.replace('_', ' ')} from the source order.")
    for field, limit in (("id",100),("customer_name",150),("assigned_employee",100)):
        if len(row.get(field, "")) > limit:
            error(field, f"Value is longer than {limit} characters.", "Shorten the value without losing its identity.")
    status = row.get("status", "").lower().replace("-", "_").replace(" ", "_")
    status = {"pending": "unconfirmed", "pending_confirmation": "unconfirmed", "shipped": "dispatched",
              "delivery_failed": "failed_delivery", "failed_delivery_attempt": "failed_delivery"}.get(status, status)
    row["status"] = status
    if status and status not in STATUSES:
        error("status", "Unknown order status; rules cannot be evaluated safely.", "Use: " + ", ".join(sorted(STATUSES)))
    payment = row.get("payment_type", "").lower()
    if payment and payment not in {"prepaid", "cod"}:
        error("payment_type", "Payment type must be Prepaid or COD.", "Replace with Prepaid or COD.")
    row["payment_type"] = "COD" if payment == "cod" else "Prepaid"
    phone = row.get("phone", "")
    digits = re.sub(r"[\s()+-]", "", phone)
    if phone and (not digits.isdigit() or not 10 <= len(digits) <= 15):
        error("phone", "Phone must contain 10–15 digits, optionally with a country prefix.", "Use the customer's complete phone number as text.")
    if digits.isdigit():
        row["phone"] = digits
    value_paise = None
    if row.get("order_value"):
        try:
            amount = Decimal(row["order_value"].replace("₹", "").replace(",", ""))
            if not amount.is_finite() or amount < 0 or amount > Decimal("1000000000") or amount * 100 != (amount * 100).to_integral_value():
                raise InvalidOperation
            value_paise = int(amount * 100)
        except (InvalidOperation, ValueError):
            error("order_value", "Order value must be a finite, nonnegative INR amount with at most two decimal places.", "Supply the actual order amount, e.g. 5000.01.")
    row["value_paise"] = value_paise
    needs_expected = status in SHIPPED or (status not in TERMINAL and value_paise is not None and value_paise > round(config["high_value_rupees"] * 100))
    for field in ("order_date", "dispatch_date", "expected_delivery_date", "confirmed_at", "delivery_attempt_at"):
        value = row.get(field, "")
        required = (field == "dispatch_date" and status in SHIPPED) or (field == "expected_delivery_date" and needs_expected)
        if not value:
            row[field] = None
            if required:
                error(field, "Date is required for this order's current status or value.", f"Supply the actual {field.replace('_', ' ')}; do not guess.")
            continue
        require_time = field in OPTIONAL or (field == "order_date" and (status == "unconfirmed" or (status == "confirmed" and not row.get("confirmed_at"))))
        try:
            row[field] = normalize_date(value, require_time)
            if field in {"order_date", "confirmed_at", "dispatch_date", "delivery_attempt_at"}:
                future = local_day(row[field]) > now.astimezone(IST).date().isoformat() if DATE_ONLY.fullmatch(row[field]) else timestamp(row[field]) > now
                if future:
                    error(field, "An event already reported as occurring cannot be in the future.", "Correct the event time or the order status.")
        except ValueError as exc:
            error(field, str(exc), "Supply a valid ISO date or timestamp from the source; hourly rules need a time.")
    if status in {"unconfirmed", "confirmed"} and row.get("dispatch_date"):
        error("dispatch_date", "An unshipped status conflicts with a dispatch date.", "Clear the date if unshipped, or change the status to dispatched.")
    # Date-order checks use known precision and do not manufacture clock times.
    if not errors:
        for earlier, later in (("order_date", "confirmed_at"), ("order_date", "dispatch_date"), ("order_date", "expected_delivery_date"), ("confirmed_at", "dispatch_date"), ("dispatch_date", "expected_delivery_date"), ("dispatch_date", "delivery_attempt_at")):
            a, b = row.get(earlier), row.get(later)
            if a and b:
                reversed_dates = local_day(b) < local_day(a) if DATE_ONLY.fullmatch(a) or DATE_ONLY.fullmatch(b) else timestamp(b) < timestamp(a)
                if reversed_dates:
                    error(later, f"{later} precedes {earlier}.", "Correct the chronology using the actual source dates.")
    return row, errors
